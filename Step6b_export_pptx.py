#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
export_pptx — 把海豚配对总览导出为 PowerPoint (单页)

export_dolphin_pptx(dolphins, out_path): 全部海豚放在同一页 16:9 幻灯片的
网格里, 每格上 DL 下 DR 代表鳍图 (先居中裁成 16:9), 两图无缝相接,
DolphinID 直接叠加在接缝处 (数据来自 ds.dolphins(), 与 Dolphin 页
Overview 模式同一份)。
"""

import io
import os

from PIL import Image
from pptx import Presentation
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

RATIO = 16 / 9  # 背鳍图统一裁成的宽高比
GRID_COLS = 5   # 固定网格列数，行数随海豚数量自动增加

def export_dolphin_pptx(dolphins, out_path, cols=GRID_COLS):
    """dolphins: [(name, dl_fid, dr_fid, dl_path, dr_path)] (Dataset.dolphins());
    所有海豚排在同一页, 每格仅叠加标注 DolphinID, 返回写入的海豚数。
    cols: 固定列数, 行数随海豚数量自动计算。"""
    prs = Presentation()
    prs.slide_width = Inches(13.333)   # 16:9
    prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # 空白版式
    os.makedirs(os.path.dirname(os.path.abspath(out_path)), exist_ok=True)
    n = len(dolphins)
    if n == 0:
        prs.save(out_path)
        return 0
    cols = max(1, int(cols))
    rows = (n + cols - 1) // cols
    W, H = 13.333, 7.5
    #cw, ch = W / cols, H / rows
    cw = W / cols
    ch = cw/16*18
    for k, (name, dl_fid, dr_fid, dl, dr) in enumerate(dolphins):
        r, c = divmod(k, cols)
        x, y = c * cw, r * ch 
        # 上 DL 下 DR 无缝相接 (各占比一半), 编号叠加在接缝处, 不占空间
        half = ch / 2
        _add_pic(slide, dl, x, y, cw, half)
        _add_pic(slide, dr, x, y + half, cw, half)
        tb = slide.shapes.add_textbox(Inches(x), Inches(y + half - 0.3),
                                      Inches(cw), Inches(0.3))
        tf = tb.text_frame
        tf.text = str(name)
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        p.runs[0].font.size = Pt(min(24, max(9, ch * 7)))
        p.runs[0].font.bold = True
    prs.save(out_path)
    return n


def _crop_to_ratio(path, ratio):
    """按图中心裁成指定宽高比, 返回 JPEG 字节流 (供 add_picture 使用)。"""
    im = Image.open(path).convert("RGB")
    w, h = im.size
    if w / h > ratio:            # 太宽 -> 左右各裁掉一部分
        nw = int(h * ratio)
        box = ((w - nw) // 2, 0, (w - nw) // 2 + nw, h)
    else:                        # 太高 -> 上下各裁掉一部分
        nh = int(w / ratio)
        box = (0, (h - nh) // 2, w, (h - nh) // 2 + nh)
    im = im.crop(box)
    buf = io.BytesIO()
    im.save(buf, "JPEG")
    buf.seek(0)
    return buf


def _add_pic(slide, path, cell_x, cell_y, cell_w, cell_h):
    """按目标格子宽高比居中裁切后填满格子; 缺图留空。"""
    if not (path and os.path.isfile(path)):
        return
    ratio = cell_w / cell_h if cell_h else RATIO
    slide.shapes.add_picture(_crop_to_ratio(path, ratio),
                             Inches(cell_x), Inches(cell_y),
                             width=Inches(cell_w), height=Inches(cell_h))


if __name__ == "__main__":
    import sys
    from doris_dataset import Dataset
    root = sys.argv[1] if len(sys.argv) > 1 else "."
    out = sys.argv[2] if len(sys.argv) > 2 \
        else os.path.join(root, "EXPORT", "dolphin_overview.pptx")
    cols = int(sys.argv[3]) if len(sys.argv) > 3 else GRID_COLS
    n = export_dolphin_pptx(Dataset(root).dolphins(), out, cols=cols)
    print("exported %d dolphins (one slide, %d cols) -> %s" % (n, cols, out))
