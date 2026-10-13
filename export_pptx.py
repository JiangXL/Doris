#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
export_pptx — 把海豚配对总览导出为 PowerPoint (单页)

export_dolphin_pptx(dolphins, out_path): 全部海豚放在同一页 16:9 幻灯片的
网格里, 每格上 DL 下 DR 代表鳍图 (先居中裁成 16:9), 两图无缝相接,
DolphinID 直接叠加在接缝处 (数据来自 ds.dolphins(), 与 Dolphin 页
Overview 模式同一份)。
"""

import io
import math
import os

from PIL import Image
from pptx import Presentation
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

RATIO = 16 / 9  # 背鳍图统一裁成的宽高比


def export_dolphin_pptx(dolphins, out_path):
    """dolphins: [(name, dl_fid, dr_fid, dl_path, dr_path)] (Dataset.dolphins());
    所有海豚排在同一页, 每格仅叠加标注 DolphinID, 返回写入的海豚数。"""
    prs = Presentation()
    prs.slide_width = Inches(13.333)   # 16:9
    prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # 空白版式
    os.makedirs(os.path.dirname(os.path.abspath(out_path)), exist_ok=True)
    n = len(dolphins)
    if n == 0:
        prs.save(out_path)
        return 0
    # 网格: 列数按 16:9 折算, 行数随之 (与 Dolphin 页 Overview 网格一致)
    cols = max(1, round(math.sqrt(n * 16 / 9)))
    #cols = 5
    rows = (n + cols - 1) // cols
    W, H = 13.333, 7.5
    cw, ch = W / cols, H / rows
    for k, (name, dl_fid, dr_fid, dl, dr) in enumerate(dolphins):
        r, c = divmod(k, cols)
        x, y = c * cw, r * ch 
        # 上 DL 下 DR 无缝相接 (各占比一半), 编号叠加在接缝处, 不占空间
        half = ch / 2
        _add_pic(slide, dl, x, y, cw, half)
        _add_pic(slide, dr, x, y + half, cw, half)
        tb = slide.shapes.add_textbox(Inches(x), Inches(y + half - 0.3),
                                      Inches(cw), Inches(0.3))
        tf = tb.text_frame
        tf.text = str(name)
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        #p.runs[0].font.size = Pt(min(16, max(9, ch * 7)))
        p.runs[0].font.size = Pt(24)
        p.runs[0].font.bold = True
    prs.save(out_path)
    return n


def _crop_169(path):
    """按图中心裁成 16:9, 返回 JPEG 字节流 (供 add_picture 使用)。"""
    im = Image.open(path).convert("RGB")
    w, h = im.size
    if w / h > RATIO:            # 太宽 -> 左右各裁掉一部分
        nw = int(h * RATIO)
        box = ((w - nw) // 2, 0, (w - nw) // 2 + nw, h)
    else:                        # 太高 -> 上下各裁掉一部分
        nh = int(w / RATIO)
        box = (0, (h - nh) // 2, w, (h - nh) // 2 + nh)
    im = im.crop(box)
    buf = io.BytesIO()
    im.save(buf, "JPEG")
    buf.seek(0)
    return buf


def _add_pic(slide, path, cell_x, cell_y, cell_w, cell_h):
    """16:9 裁切后放入格子 (占满给定高度, 水平居中); 缺图留空。"""
    if not (path and os.path.isfile(path)):
        return
    w_in = min(cell_h * RATIO, cell_w)   # 裁后恰为 16:9, 高随宽定
    h_in = w_in / RATIO
    left = Inches(cell_x + (cell_w - w_in) / 2)
    slide.shapes.add_picture(_crop_169(path), left, Inches(cell_y),
                             width=Inches(w_in), height=Inches(h_in))


if __name__ == "__main__":
    import sys
    from doris_dataset import Dataset
    root = sys.argv[1] if len(sys.argv) > 1 else "."
    out = sys.argv[2] if len(sys.argv) > 2 \
        else os.path.join(root, "EXPORT", "dolphin_overview.pptx")
    n = export_dolphin_pptx(Dataset(root).dolphins(), out)
    print("exported %d dolphins (one slide) -> %s" % (n, out))
